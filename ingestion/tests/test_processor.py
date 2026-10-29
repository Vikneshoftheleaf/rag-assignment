import tempfile
import unittest
from pathlib import Path
import sys
from types import SimpleNamespace
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from processor import chunk_text, chunks_from_docling, document_type, markdown_fallback_chunks


class ProcessorTests(unittest.TestCase):
    def test_pdfium_fallback_extracts_text_with_page_provenance(self):
        from processor import pdfium_fallback_chunks

        class FakeTextPage:
            def __init__(self, text: str) -> None:
                self.text = text

            def get_text_range(self) -> str:
                return self.text

            def close(self) -> None:
                pass

        class FakePage:
            def __init__(self, text: str) -> None:
                self.text = text

            def get_textpage(self) -> FakeTextPage:
                return FakeTextPage(self.text)

            def close(self) -> None:
                pass

        class FakeDocument(list):
            def close(self) -> None:
                pass

        document = FakeDocument([FakePage("Page one text."), FakePage("Page two text.")])
        pdfium = SimpleNamespace(PdfDocument=lambda _path: document)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.pdf"
            path.touch()
            with patch.dict(sys.modules, {"pypdfium2": pdfium}):
                chunks = pdfium_fallback_chunks(path)

        self.assertEqual([chunk["content"] for chunk in chunks], ["Page one text.", "Page two text."])
        self.assertEqual(
            [chunk["source_location"] for chunk in chunks],
            [{"page": 1}, {"page": 2}],
        )

    def test_heading_and_pdf_page_are_preserved(self):
        doc = {
            "texts": [
                {"label": "section_header", "text": "Results", "prov": [{"page_no": 4}]},
                {"label": "text", "text": "Revenue increased.\n\nOperating costs were stable.", "prov": [{"page_no": 4}]},
            ],
            "tables": [],
        }
        chunks = chunks_from_docling(doc)
        self.assertEqual(len(chunks), 1)
        self.assertIn("Results", chunks[0]["content"])
        self.assertEqual(chunks[0]["source_location"]["page"], 4)
        self.assertEqual(chunks[0]["metadata"]["section"], "Results")

    def test_table_rows_retain_header_and_location(self):
        doc = {
            "texts": [],
            "tables": [{
                "prov": [{"page_no": 2}],
                "data": {"table_cells": [
                    {"start_row_offset_idx": 0, "start_col_offset_idx": 0, "text": "Metric"},
                    {"start_row_offset_idx": 0, "start_col_offset_idx": 1, "text": "H1"},
                    {"start_row_offset_idx": 0, "start_col_offset_idx": 2, "text": "H2"},
                    {"start_row_offset_idx": 1, "start_col_offset_idx": 0, "text": "Revenue"},
                    {"start_row_offset_idx": 1, "start_col_offset_idx": 2, "text": "100"},
                ]},
            }],
        }
        chunks = chunks_from_docling(doc)
        self.assertIn("Metric | H1 | H2", chunks[0]["content"])
        self.assertIn("Revenue |  | 100", chunks[0]["content"])
        self.assertEqual(chunks[0]["source_location"]["page"], 2)
        self.assertEqual(chunks[0]["source_location"]["table"], 1)

    def test_long_text_chunks_without_losing_source_metadata(self):
        chunks = chunk_text("\n\n".join(["Evidence text."] * 140), heading="Minutes", location={"section": "Minutes"})
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(item["source_location"]["section"] == "Minutes" for item in chunks))

    def test_oversized_paragraph_is_split_without_dropping_or_duplicating_text(self):
        source = "Revenue increased in Q2. " * 179 + "Revenue increased in Q2."
        chunks = chunk_text(source)

        self.assertGreater(len(chunks), 1)
        self.assertEqual("".join(chunk["content"] for chunk in chunks), source)
        self.assertTrue(all(len(chunk["content"]) <= 1100 for chunk in chunks))

    def test_paragraph_boundaries_are_preserved_across_chunks(self):
        source = "\n\n".join(["Revenue increased."] * 100)
        chunks = chunk_text(source)

        self.assertEqual("".join(chunk["content"] for chunk in chunks), source)

    def test_markdown_fallback_tracks_heading(self):
        chunks = markdown_fallback_chunks("# Board outcome\n\nThe board approved the report.")
        self.assertIn("Board outcome", chunks[0]["content"])
        self.assertEqual(chunks[0]["metadata"]["section"], "Board outcome")

    def test_extension_validation(self):
        self.assertEqual(document_type(Path("report.PDF")), "pdf")
        with self.assertRaises(ValueError):
            document_type(Path("notes.txt"))

    def test_spreadsheet_facts_keep_row_cell_period_scope_and_units(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Consolidated"
            sheet.append(["Metric", "H1 FY2026-27 (₹ Crore)"])
            sheet.append(["Revenue", 125.5])
            sheet.append(["EBITDA", "=SUM(B2:B2)"])
            workbook.save(path)

            facts, warnings = spreadsheet_facts(path)
            self.assertEqual(len(facts), 2)
            self.assertEqual(facts[0]["metric"], "Revenue — H1 FY2026-27 (₹ Crore)")
            self.assertEqual(facts[0]["scope"], "consolidated")
            self.assertEqual(facts[0]["period"], "H1 FY2026-27")
            self.assertEqual(facts[0]["unit"], "₹ crore")
            self.assertEqual(facts[0]["numeric_value"], 125.5)
            self.assertEqual(facts[0]["source_location"], {"sheet": "Consolidated", "row": 2, "cell": "B2"})
            self.assertEqual(facts[1]["metric"], "EBITDA — H1 FY2026-27 (₹ Crore)")
            self.assertEqual(facts[1]["numeric_value"], 125.5)
            self.assertTrue(facts[1]["metadata"]["formula_evaluated"])
            self.assertFalse(warnings)

    def test_as_at_dates_override_filename_reporting_period(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results_Q2_H1_FY2026-27.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Balance Sheet"
            sheet.append(["Particulars", "As at 30-Sep-2026", "As at 31-Mar-2026"])
            sheet.append(["Property, plant and equipment", 680, 640])
            workbook.save(path)

            facts, _ = spreadsheet_facts(path)

        self.assertEqual(
            [fact["period"] for fact in facts],
            ["As at 30-Sep-2026", "As at 31-Mar-2026"],
        )
        self.assertEqual([fact["numeric_value"] for fact in facts], [680, 640])

    def test_formula_values_use_evaluated_result_not_cached_cell_value(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Standalone Results"
            sheet.append(["Particulars", "Q2 FY2026-27"])
            sheet.append(["Revenue", 120])
            sheet.append(["Total income", "=B2+5"])
            workbook.save(path)

            facts, warnings = spreadsheet_facts(path)

        formula_fact = next(fact for fact in facts if fact["metadata"]["row_label"] == "Total income")
        self.assertEqual(formula_fact["numeric_value"], 125)
        self.assertEqual(formula_fact["value_text"], "125")
        self.assertTrue(formula_fact["metadata"]["formula_evaluated"])
        self.assertFalse(warnings)

    def test_iferror_formulas_evaluate_values_and_fallbacks_in_workbook_cells(self):
        from openpyxl import Workbook
        from processor import spreadsheet_cells

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            results = workbook.active
            results.title = "Standalone Results"
            results.append(["Particulars", "Q2 FY2026-27"])
            results.append(["Revenue", 120])
            dashboard = workbook.create_sheet("Dashboard")
            dashboard.append(["Metric", "Value"])
            dashboard.append(["Growth", "='Standalone Results'!B2/0"])
            dashboard["B2"] = "=IFERROR('Standalone Results'!B2/0,0)"
            dashboard["C2"] = "='Standalone Results'!B10"
            dashboard.append(["Margin", "=IFERROR('Standalone Results'!B2/2,0)"])
            dashboard.append(["Difference", "=B2-B3"])
            workbook.save(path)

            cells, warnings = spreadsheet_cells(path)

        indexed = {cell["cell_address"]: cell for cell in cells}
        self.assertEqual(indexed["B2"]["value_text"], "0")
        self.assertEqual(indexed["C2"]["value_text"], "0")
        self.assertEqual(indexed["B3"]["value_text"], "60")
        self.assertEqual(indexed["B4"]["value_text"], "-60")
        self.assertTrue(indexed["B2"]["formula_evaluated"])
        self.assertTrue(indexed["C2"]["formula_evaluated"])
        self.assertTrue(indexed["B3"]["formula_evaluated"])
        self.assertTrue(indexed["B4"]["formula_evaluated"])
        self.assertFalse(warnings)

    def test_spreadsheet_manifest_accounts_for_every_worksheet(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts, spreadsheet_manifest

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            workbook.active.title = "Results"
            workbook.active.append(["Particulars", "Q2 FY2026-27"])
            workbook.active.append(["Revenue", 120])
            workbook.create_sheet("Narrative").append(["Notes"])
            workbook.create_sheet("Dashboard").append(["Summary"])
            workbook.save(path)

            facts, _ = spreadsheet_facts(path)
            manifest = spreadsheet_manifest(path, facts)

        self.assertEqual([sheet["sheet"] for sheet in manifest], ["Results", "Narrative", "Dashboard"])
        self.assertEqual(manifest[0]["disposition"], "structured_facts")
        self.assertEqual(manifest[0]["structured_fact_count"], 1)
        self.assertEqual(manifest[1]["disposition"], "narrative_or_unstructured")
        self.assertEqual(manifest[2]["disposition"], "dashboard_cells")

    def test_spreadsheet_chunks_preserve_every_sheet_and_cell_address(self):
        from openpyxl import Workbook
        from processor import spreadsheet_cells, spreadsheet_chunks

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "workbook.xlsx"
            workbook = Workbook()
            results = workbook.active
            results.title = "Results"
            results.append(["Particulars", "Q2 FY2026-27"])
            results.append(["Revenue from operations", 1245.6])
            notes = workbook.create_sheet("Notes")
            notes["A1"] = "Review status"
            notes["B1"] = "Pending final approval"
            notes.merge_cells("A3:B3")
            notes["A3"] = "Narrative across merged cells"
            workbook.save(path)

            chunks = spreadsheet_chunks(path)
            cells, warnings = spreadsheet_cells(path)

        content = "\n".join(chunk["content"] for chunk in chunks)
        self.assertIn("Sheet: Results", content)
        self.assertIn(
            '{"type":"workbook_cell","sheet":"Results","row":2,'
            '"row_label":"Revenue from operations","column":"B",'
            '"column_header":"Q2 FY2026-27","cell":"B2","value":"1245.6"}',
            content,
        )
        self.assertIn("Sheet: Notes", content)
        self.assertIn("Pending final approval", content)
        self.assertIn("Narrative across merged cells", content)
        revenue_chunk = next(
            chunk for chunk in chunks if '"cell":"B2","value":"1245.6"' in chunk["content"]
        )
        self.assertTrue(revenue_chunk["content"].startswith("Sheet: Results"))
        self.assertEqual(revenue_chunk["source_location"]["sheet"], "Results")
        self.assertEqual(revenue_chunk["metadata"]["sheet"], "Results")
        self.assertEqual(revenue_chunk["source_location"]["row"], 2)
        self.assertEqual(revenue_chunk["source_location"]["cell"], "A2")
        revenue_value = next(cell for cell in cells if cell["cell_address"] == "B2")
        self.assertEqual(revenue_value["row_label"], "Revenue from operations")
        self.assertEqual(revenue_value["column_header"], "Q2 FY2026-27")
        self.assertEqual(revenue_value["numeric_value"], 1245.6)
        self.assertFalse(warnings)

    def test_spreadsheet_headers_and_vertical_form_fields_keep_distinct_labels(self):
        from openpyxl import Workbook
        from processor import spreadsheet_cells, spreadsheet_chunks

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ctc.xlsx"
            workbook = Workbook()
            register = workbook.active
            register.title = "CTC Register"
            register.append(["CTC ID", "Board Meeting Date", "Resolution No."])
            register.append(["SIL-CTC-2026-001", "22-Oct-2026", "BR-01"])
            form = workbook.create_sheet("BR-01 Results")
            form.append([
                "CTC ID", "SIL-CTC-2026-001", "Resolution No.", "BR-01",
                "Date of Board Meeting", "22 October 2026",
            ])
            workbook.save(path)

            cells, warnings = spreadsheet_cells(path)
            chunks = spreadsheet_chunks(path, cells)

        by_address = {cell["cell_address"]: cell for cell in cells}
        self.assertEqual(by_address["B2"]["column_header"], "Board Meeting Date")
        self.assertEqual(by_address["D1"]["row_label"], "Resolution No.")
        self.assertEqual(by_address["D1"]["column_header"], "Value")
        self.assertIn(
            '{"type":"workbook_cell","sheet":"BR-01 Results","row":1,'
            '"row_label":"Resolution No.","column":"D","column_header":"Value",'
            '"cell":"D1","value":"BR-01"}',
            next(chunk["content"] for chunk in chunks if '"cell":"D1"' in chunk["content"]),
        )
        self.assertFalse(warnings)

    def test_financial_facts_keep_quarter_comparatives_and_inherited_units(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results_Q2_H1_FY2026-27.xlsx"
            workbook = Workbook()
            standalone = workbook.active
            standalone.title = "Standalone Results"
            consolidated = workbook.create_sheet("Consolidated Results")

            for sheet, values in (
                (standalone, [1245.6, 1108.2, 2430.8, 2165.4]),
                (consolidated, [1378.4, 1216.7, 2690.5, 2381.2]),
            ):
                sheet["A1"] = "UNAUDITED FINANCIAL RESULTS"
                sheet["A2"] = "₹ crore except EPS"
                sheet.append([])
                sheet.append([])
                sheet.append([])
                sheet.append([
                    "Particulars",
                    "Quarter ended 30-Sep-2026",
                    "Quarter ended 30-Sep-2025",
                    "Half year ended 30-Sep-2026",
                    "Half year ended 30-Sep-2025",
                ])
                sheet.append(["Revenue from operations", *values])
                sheet.append(["Basic EPS", 12.5, 11.1, 23.4, 21.6])

            workbook.save(path)
            facts, _ = spreadsheet_facts(path)

        revenue_facts = [fact for fact in facts if "Revenue from operations" in fact["metric"]]
        eps_facts = [fact for fact in facts if "Basic EPS" in fact["metric"]]
        self.assertEqual(len(revenue_facts), 8)
        self.assertEqual(len(eps_facts), 8)
        self.assertTrue(all(fact["unit"] == "₹ per share" for fact in eps_facts))
        self.assertEqual(
            [(fact["period"], fact["unit"], fact["value_text"]) for fact in revenue_facts[:4]],
            [
                ("Q2 FY2026-27", "₹ crore", "1245.6"),
                ("Q2 FY2025-26", "₹ crore", "1108.2"),
                ("H1 FY2026-27", "₹ crore", "2430.8"),
                ("H1 FY2025-26", "₹ crore", "2165.4"),
            ],
        )
        self.assertEqual(
            [fact["scope"] for fact in revenue_facts],
            ["standalone"] * 4 + ["consolidated"] * 4,
        )
        self.assertEqual(
            [fact["numeric_value"] for fact in revenue_facts[4:]],
            [1378.4, 1216.7, 2690.5, 2381.2],
        )

    def test_formula_cells_are_evaluated_and_indexed_as_numeric_facts(self):
        from openpyxl import Workbook
        from processor import spreadsheet_cells, spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Standalone Results"
            sheet.append(["Particulars", "Q2 FY2026-27"])
            sheet.append(["Revenue from operations", 100])
            sheet.append(["Other income", 5])
            sheet.append(["Total income", "=B2+B3"])
            sheet.append(["Profit for the period", "=SUM(B2:B3)-10"])
            workbook.save(path)

            facts, warnings = spreadsheet_facts(path)
            cells, cell_warnings = spreadsheet_cells(path)

        computed = {fact["metadata"]["row_label"]: fact for fact in facts}
        computed_cells = {cell["row_label"]: cell for cell in cells if cell["column_number"] == 2}
        self.assertEqual(computed["Total income"]["numeric_value"], 105)
        self.assertEqual(computed["Profit for the period"]["numeric_value"], 95)
        self.assertEqual(computed["Profit for the period"]["value_text"], "95")
        self.assertTrue(computed["Profit for the period"]["metadata"]["formula_evaluated"])
        self.assertEqual(computed_cells["Total income"]["numeric_value"], 105)
        self.assertEqual(computed_cells["Profit for the period"]["value_text"], "95")
        self.assertEqual(computed_cells["Profit for the period"]["column_header"], "Q2 FY2026-27")
        self.assertTrue(computed_cells["Profit for the period"]["formula_evaluated"])
        self.assertFalse(warnings)
        self.assertFalse(cell_warnings)

    def test_unsupported_formula_is_kept_unresolved_instead_of_indexing_none_as_value(self):
        from openpyxl import Workbook
        from processor import spreadsheet_cells

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Particulars", "Q2 FY2026-27"])
            sheet.append(["Unsupported result", '=A2&" text"'])
            workbook.save(path)

            cells, warnings = spreadsheet_cells(path)

        formula_cell = next(cell for cell in cells if cell["cell_address"] == "B2")
        self.assertEqual(formula_cell["cell_type"], "formula")
        self.assertEqual(formula_cell["value_text"], "")
        self.assertIsNone(formula_cell["numeric_value"])
        self.assertIn("B2: formula result is unavailable or unsupported", warnings[0])

    def test_formula_evaluator_resolves_cross_sheet_references(self):
        from openpyxl import Workbook
        from processor import spreadsheet_facts

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            workbook = Workbook()
            results = workbook.active
            results.title = "Standalone Results"
            results.append(["Particulars", "Q2 FY2026-27"])
            results.append(["Revenue", 120])
            ratios = workbook.create_sheet("EPS & Ratios")
            ratios.append(["Particulars", "Q2 FY2026-27"])
            ratios.append(["PAT", 80])
            ratios.append(["Basic EPS", "=B2/10"])
            results.append(["Basic EPS", "='EPS & Ratios'!B3"])
            workbook.save(path)

            facts, warnings = spreadsheet_facts(path)

        eps_facts = [fact for fact in facts if fact["metadata"]["row_label"] == "Basic EPS"]
        self.assertEqual([fact["numeric_value"] for fact in eps_facts], [8, 8])
        self.assertTrue(all(fact["metadata"]["formula_evaluated"] for fact in eps_facts))
        self.assertFalse(warnings)


if __name__ == "__main__":
    unittest.main()
